"""Recalculate an .xlsx in Microsoft Excel and save it, then report formula errors.

openpyxl writes formulas without computed values, so anything that reads cached
values (pandas, the pipeline) sees blanks until Excel has calculated and saved
the file. Windows + Excel + pywin32 only.

    python scripts/recalc_excel.py workbook/valuation.xlsx
"""
import sys
from pathlib import Path

import win32com.client
from openpyxl import load_workbook

ERRORS = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#NULL!")


def recalc(path: Path) -> None:
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(path.resolve()))
        excel.CalculateFull()
        wb.Save()
        wb.Close(SaveChanges=False)
    finally:
        excel.Quit()


def find_errors(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True)
    return [f"{ws.title}!{cell.coordinate}: {cell.value}"
            for ws in wb.worksheets for row in ws.iter_rows() for cell in row
            if isinstance(cell.value, str) and cell.value in ERRORS]


if __name__ == "__main__":
    target = Path(sys.argv[1])
    recalc(target)
    errors = find_errors(target)
    print(f"Recalculated {target}: {len(errors)} formula errors")
    for e in errors:
        print("  ", e)
    sys.exit(1 if errors else 0)
